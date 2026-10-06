"""Rebuild the HRI study design form for the ESP32 screen version.

Takes the block-handover form (same course layout) and replaces the answer
text under each prompt. Headings, prompts and the grey Hoffman and Zhao hint
lines are kept as they are.

Usage:
    python tools/build_form_esp32.py docs/source/HRI_Study_Design_Final_Presentable_blocks.docx \
        docs/HRI_Study_Design_ESP32.docx
"""
import re
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Line prefixes: "LIST:" bullet, "KN:" keep with next paragraph, "KNLIST:" both,
# "" blank line.
# **text** is bold.
BLOCKS = {
    # More specific research questions (paragraphs 7 to 11)
    (7, 11): [
        "How does coordinated versus uncoordinated timing of a Kinova Gen3 arm that shows "
        "information on a small screen affect:",
        "(1) subjective workload and disturbance,",
        "(2) how correctly and completely participants register the information,",
        "(3) compensatory behaviour such as slower responses, hesitation, or extra effort, and",
        "(4) which setup participants would choose to keep working with?",
    ],
    # Constructs (16 to 27)
    (16, 27): [
        "The study looks at one predictor construct and four outcome constructs. NASA-TLX is "
        "used as a measure of workload; it is not the construct itself.",
        "",
        "KN:**Predictor construct**",
        "LIST:**Robot temporal coordination:** how well the robot's timing fits the "
        "participant: whether it holds the screen still long enough to read and moves at "
        "regular, predictable moments. We manipulate this with a coordinated and an "
        "uncoordinated condition.",
        "",
        "KN:**Outcome constructs**",
        "LIST:**Subjective workload / disturbance:** how mentally demanding, rushed, "
        "effortful, frustrating, or uncomfortable the task feels.",
        "LIST:**Objective task performance:** how correctly and completely the participant "
        "registers the colour and number shown on the screen.",
        "LIST:**Compensatory behaviour:** changes in attention or response behaviour that "
        "help the participant keep performing when the robot timing becomes harder to "
        "manage. We look for this mainly in response time, hesitation, and reported effort.",
        "LIST:**Robot acceptance:** whether the participant would choose to keep working with "
        "the robot in a given timing condition.",
        "",
        "We do not add a secondary task such as a running sum. Each item has to be held in "
        "memory until it is entered, and in the uncoordinated condition a new item can arrive "
        "before the previous one is entered, so the task already puts load on memory and "
        "attention. A second task would add a factor our sample size cannot handle.",
    ],
    # Hypotheses (32 to 39)
    (32, 39): [
        "The coordinated condition is the baseline for all comparisons. Before testing the "
        "main hypotheses, we check that participants actually perceive the uncoordinated "
        "condition as less predictable and as giving them less time to read. This is a "
        "manipulation check rather than a separate hypothesis.",
        "",
        "LIST:**H1. Subjective workload:** Participants will report higher subjective workload "
        "in the uncoordinated condition than in the coordinated condition.",
        "LIST:**H2. Compensatory behaviour:** In the uncoordinated condition, participants will "
        "respond later and less regularly (longer and more variable response time), will "
        "hesitate more, and will report more effort.",
        "LIST:**H3. Task performance:** The items participants enter will be about as accurate "
        "in both conditions, because entering the items correctly is the task we tell "
        "participants to focus on. Missed items will increase in the uncoordinated condition.",
        "LIST:**H4. Robot acceptance:** When asked which setup they would use for a full work "
        "shift, most participants will choose the coordinated one. We expect this even though "
        "the accuracy of the entered items stays the same (H3), because the choice should "
        "follow how the task felt, not how well it went.",
        "",
        "Together, these answer our research question. If workload rises (H1) but the entered "
        "items stay accurate (H3), the cost has gone somewhere else: into slower and less "
        "regular responses (H2), into missed items (H3), and into which robot behaviour people "
        "are willing to work with (H4).",
    ],
    # Experimental conditions (44 to 52)
    (44, 52): [
        "The independent variable is robot temporal coordination, with two conditions. In "
        "both, one Kinova Gen3 arm holds an ESP32 board with a small screen. The arm brings the "
        "screen to a fixed viewing position in front of the participant and holds it still "
        "while the screen shows one item. The screen then goes blank, the arm moves it back to "
        "a fixed waiting position, and the next item is loaded. Each movement takes about 1 "
        "second.",
        "LIST:**Coordinated condition (baseline):** The robot follows a fixed rhythm. A new "
        "item appears every 5 seconds and the screen always stays still for 2 seconds, so the "
        "participant has enough time to read each item and can anticipate the next one.",
        "LIST:**Uncoordinated condition:** The averages stay the same (one item every 5 "
        "seconds, 2 seconds of viewing time), but the timing varies. The viewing time changes "
        "between 1 and 3 seconds, so the screen sometimes leaves early. The time from one item "
        "to the next changes between 3 and 7 seconds, so the next item sometimes arrives "
        "before the participant has finished entering the last one. The timing sequence is "
        "prepared in advance so that the averages are the same as in the coordinated "
        "condition.",
        "",
        "Even the shortest viewing time is long enough to read an item when the participant is "
        "looking at the screen. A missed item therefore shows that attention was somewhere "
        "else, not that the item could not be read.",
        "",
        "Both conditions use the same arm path, speed and acceleration, viewing and waiting "
        "positions, screen, item format, number of items, trial length, and instructions. "
        "Only the timing pattern changes.",
        "",
        "Each item is a colour word (red, blue, green or yellow) and a number from 1 to 9, for "
        "example “RED 4”. The colour word is printed in its own colour. The "
        "participant reads the item and enters it on a tablet in front of them by tapping one "
        "colour button and one number button. If they did not catch an item, they tap "
        "“Missed” and continue. Entering each item correctly is explicitly described "
        "as the main task, and participants are asked to enter each item as soon as they have "
        "read it. Each trial contains 20 items and lasts about 2 minutes, so participants work "
        "under time pressure in both conditions.",
        "",
        "We will use two matched item lists with the same number of items and the same colour "
        "and number distribution, but in a different order. The timing range will be checked "
        "in a short pilot before the main study and then kept fixed for all participants.",
    ],
    # Ordering effects (61)
    (61, 61): [
        "The order will be counterbalanced. Half of the participants do the coordinated "
        "condition first and the uncoordinated condition second; the other half do the "
        "reverse. Participants are assigned to one of the two orders before the session. "
        "Everyone completes a 1-minute practice trial with 10 items first. The practice mixes "
        "regular and irregular timing, so it does not favour either condition. Between the two "
        "experimental trials, the participant fills in the questionnaires while the "
        "experimenter loads the next trial. The matched item lists are also swapped across "
        "conditions, so one condition is not always paired with the same list. Crossing the "
        "two orders with the two lists gives four groups of four participants.",
    ],
    # Manipulation check (67 to 70)
    (67, 70): [
        "After each condition, right after NASA-TLX, participants answer two 7-point items "
        "(1 = strongly disagree, 7 = strongly agree):",
        "LIST:“The timing of the robot was predictable.”",
        "LIST:“The robot gave me enough time to read each item.”",
        "This confirms that participants noticed the timing difference. If they did not, the "
        "other results cannot be linked to the timing, so we report this check before the main "
        "results. We ask these items after NASA-TLX, so that they do not point participants to "
        "the timing before they rate their workload.",
    ],
    # Objective measures (78 to 83)
    (78, 83): [
        "The main objective measures are:",
        "LIST:**Correct registrations:** the number and percentage of the 20 items entered "
        "with both the right colour and the right number. We also report the share of entered "
        "items that are correct.",
        "LIST:**Missed items:** the number of items that were shown but not entered, including "
        "items marked as “Missed”.",
        "LIST:**Response time:** the time from the item appearing on the screen until the "
        "participant finishes entering it. We compare both the average and how variable it is "
        "across the trial. We only use items that appeared after the previous entry was "
        "finished, so that time spent on the previous item does not count.",
        "LIST:**Hesitations:** the number of entries the participant corrects with the "
        "“Undo” button.",
        "LIST:**Overlaps:** the number of times a new item appeared while the participant was "
        "still entering the previous one. Because this depends partly on the timing schedule, "
        "it is a secondary, descriptive measure.",
        "The robot computer logs when each item appears and disappears, and every tap on the "
        "tablet is logged on the same computer, so all objective measures come from the logs "
        "and no live scoring is needed. Before data collection, we name one main measure per "
        "hypothesis: Raw NASA-TLX (H1), response time (H2), missed items (H3) and the setup "
        "choice (H4). The other measures are reported as exploratory.",
    ],
    # Procedure (93 to 105)
    (93, 105): [
        "The participant sits at a table facing one Kinova Gen3 arm, with a tablet on the "
        "table in front of them. The arm holds an ESP32 board with a small screen. The consent "
        "form is signed before the session starts. After a short safety briefing, the "
        "experimenter explains the task using the same instructions for every participant. "
        "The robot will repeatedly bring the screen in front of the participant, each time "
        "showing a colour and a number. The participant must read each item and enter it on "
        "the tablet as soon as they have read it, by tapping the colour and then the number. "
        "Entering each item correctly is described as the main task. We do not tell "
        "participants which condition we expect to be more demanding. Before data collection "
        "starts, they complete a 1-minute practice trial with 10 items.",
        "",
        "The participant then completes two experimental trials of 20 items each, about 2 "
        "minutes each, one coordinated and one uncoordinated, in their counterbalanced order. "
        "Immediately after each trial, the participant completes NASA-TLX and then the "
        "manipulation-check items while the experimenter loads the next trial. The robot "
        "computer and the tablet log all objective measures automatically. After both trials, "
        "the participant answers the setup-choice question and the two interview questions, "
        "and is then debriefed about the purpose of the timing manipulation. The robot never "
        "touches the participant and never comes closer than the viewing position, about 60 cm "
        "in front of their face. A researcher stays next to the emergency stop throughout the "
        "robot task.",
        "",
        "KN:**Session timing**",
        "KNLIST:Safety briefing and instructions: 1 minute",
        "KNLIST:Practice trial (10 items): 1 minute",
        "KNLIST:Trial 1 (20 items): 2 minutes",
        "KNLIST:NASA-TLX and manipulation check: 1.5 minutes",
        "KNLIST:Trial 2 (20 items): 2 minutes",
        "KNLIST:NASA-TLX and manipulation check: 1.5 minutes",
        "KNLIST:Setup choice, interview and debrief: 1 minute",
        "LIST:**Total: about 10 minutes per participant**",
    ],
    # Sample size, last paragraph (115)
    (115, 115): [
        "Because the items are shown as text on a small screen, participants need normal or "
        "corrected-to-normal vision. The colour word can be read without colour vision, so "
        "colour vision deficiency does not exclude anyone. The practice trial confirms that "
        "each participant can read the screen from their seat. A small pilot with lab members "
        "is used to check the timing and task difficulty, and pilot data are not included in "
        "the final dataset.",
    ],
}

# Paragraphs that must still read as in the block version before we touch them.
# This guards against running the script on a different template.
EXPECTED_START = {
    7: "How does coordinated versus uncoordinated handover timing",
    16: "The study looks at one predictor construct",
    32: "The coordinated condition is the baseline",
    44: "The independent variable is robot temporal coordination",
    61: "The order will be counterbalanced.",
    67: "After each condition, participants answer two 7-point items",
    78: "The main objective measures are:",
    93: "The participant stands at a workstation",
    115: "Because the sorting task depends on colour",
}


def make_bullet(p, num_id=1):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numPr.append(ilvl)
    numPr.append(num)
    pPr.append(numPr)


def rich(p, text):  # supports **bold**
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part:
            p.add_run(part)


def add_line(anchor, line):
    p = anchor.insert_paragraph_before()
    if line.startswith("KNLIST:"):
        rich(p, line[7:])
        make_bullet(p)
        p.paragraph_format.keep_with_next = True
    elif line.startswith("LIST:"):
        rich(p, line[5:])
        make_bullet(p)
    elif line.startswith("KN:"):
        rich(p, line[3:])
        p.paragraph_format.keep_with_next = True
    else:
        rich(p, line)


def main(src, dst):
    doc = docx.Document(src)
    P = list(doc.paragraphs)  # captured once, so indices keep pointing at the originals

    for idx, start in EXPECTED_START.items():
        if not P[idx].text.startswith(start):
            sys.exit(f"Template mismatch at paragraph {idx}: {P[idx].text[:60]!r}")

    for (first, last), lines in BLOCKS.items():
        for line in lines:
            add_line(P[first], line)
        for old in P[first:last + 1]:
            old._p.getparent().remove(old._p)

    text = "\n".join(p.text for p in doc.paragraphs)
    bad = [c for c in "\u2013\u2014\u2192\u2190\u21d2" if c in text]
    if bad:
        sys.exit(f"Forbidden characters in output: {bad}")
    doc.save(dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
